"""表格文件读取与取列:CSV / Excel → 图表数据,CLI --file 与 MCP file 共用。

CSV 兼容 UTF-8(-sig)与 GBK(Excel 中文导出常见);Excel 仅支持 .xlsx / .xlsm。
首行一律视为表头;列约定:类目列默认第 1 列,数值列 = 其余列(或 col 挑 1 列),
各取数函数对应一类图表的列布局(bar/line/combo/box/xy/hist/heatmap,见各 docstring)。
"""
from __future__ import annotations

import csv
from pathlib import Path

__all__ = [
    'read_table', 'to_float', 'pick_column', 'file_columns', 'file_box_groups',
    'file_xy', 'file_samples', 'file_matrix',
]


def read_table(path: str | Path) -> tuple[list[str], list[list[str]]]:
    """读取表格,返回 (表头, 数据行);空行剔除,单元格已去空白。"""
    p = Path(path)
    suffix = p.suffix.lower()
    if suffix not in ('.csv', '.xlsx', '.xlsm'):
        raise ValueError(f'不支持的表格格式 {suffix!r}(支持 .csv / .xlsx)')
    if not p.exists():
        raise ValueError(f'文件不存在: {path}')
    header, rows = _read_csv(p) if suffix == '.csv' else _read_excel(p)
    if not header:
        raise ValueError('表格为空')
    return header, rows


def to_float(cell, where: str = '') -> float:
    """单元格转数值:容忍千分位逗号与空白;失败时报出位置便于定位。"""
    text = str(cell).replace(',', '').strip()
    try:
        return float(text)
    except ValueError as e:
        raise ValueError(f'数值列出现非数值单元格: {cell!r}{where}') from e


def pick_column(spec: str | None, header: list[str], default: int | None = None) -> int | None:
    """按表头名或从 1 数的序号定位列;spec 为空时返回 default。"""
    if spec is None:
        return default
    if spec in header:
        return header.index(spec)
    try:
        idx = int(spec) - 1
    except ValueError:
        raise ValueError(f'找不到列 {spec!r};可用列:{", ".join(header)}') from None
    if not 0 <= idx < len(header):
        raise ValueError(f'列序号 {spec} 超出范围,共 {len(header)} 列')
    return idx


def file_columns(path, cat_col=None, col=None) -> tuple[list[str], list[tuple[str, list[float]]]]:
    """类目型图表取数(bar/line/area/pie/donut/waterfall/funnel/combo/多系列):
    类目列(默认第 1 列)+ 数值列(默认除类目列外全部;col 挑 1 列),
    返回 (类目, [(列名, 数值列表)]),数值列空单元格剔除。"""
    header, rows = read_table(path)
    cat_i = pick_column(cat_col, header, 0)
    cats = [r[cat_i] if len(r) > cat_i else '' for r in rows]
    val_is = [pick_column(col, header)] if col else \
        [j for j in range(len(header)) if j != cat_i]
    cols = []
    for j in val_is:
        nm = header[j] or f'列{j + 1}'
        cols.append((nm, [to_float(r[j], f'(列「{nm}」)')
                          for r in rows if len(r) > j and r[j] != '']))
    return cats, cols


def file_box_groups(path) -> list[tuple[str, list[float]]]:
    """箱线图取数:每列一组,首行表头 = 组名,单元格 = 原始样本(各组无需对齐)。"""
    header, rows = read_table(path)
    return [(name or f'组{j + 1}',
             [to_float(r[j], f'(列「{name}」)')
              for r in rows if len(r) > j and r[j] != ''])
            for j, name in enumerate(header)]


def file_xy(path, *, sizes=False) -> tuple[list[float], list[float],
                                           list[float] | None, list[str] | None]:
    """散点/气泡取数:列顺序 x, y(, sizes)(, labels),各列空单元格独立剔除
    (可能与 x/y 不对齐,由校验兜底);sizes=True 才读第 3 列气泡大小。"""
    header, rows = read_table(path)
    xs = [to_float(r[0], '(列 1)') for r in rows if r and r[0] != '']
    ys = [to_float(r[1], '(列 2)') for r in rows if len(r) > 1 and r[1] != '']
    sz = [to_float(r[2], '(列 3)') for r in rows if len(r) > 2 and r[2] != ''] \
        if sizes else None
    labels = [r[3] for r in rows if len(r) > 3 and r[3]] or None
    return xs, ys, sz, labels


def file_samples(path) -> list[float]:
    """直方图取数:第 1 列为原始样本。"""
    header, rows = read_table(path)
    vals = [to_float(r[0], f'(列「{header[0]}」)') for r in rows if r and r[0] != '']
    if not vals:
        raise ValueError('第 1 列需为数值样本')
    return vals


def file_matrix(path, cat_col=None) -> tuple[list[str], list[str], list[list[float]]]:
    """热力图取数:类目列(默认第 1 列)= 行名,其余列为矩阵(表头 = 列名)。"""
    header, rows = read_table(path)
    cat_i = pick_column(cat_col, header, 0)
    rlabels = [r[cat_i] if len(r) > cat_i else '' for r in rows]
    val_is = [j for j in range(len(header)) if j != cat_i]
    cols = [header[j] or f'列{j + 1}' for j in val_is]
    matrix = [[to_float(r[j], f'(列「{header[j]}」)') for j in val_is] for r in rows]
    return rlabels, cols, matrix


def _read_csv(path) -> tuple[list[str], list[list[str]]]:
    raw = Path(path).read_bytes()
    for enc in ('utf-8-sig', 'gbk'):
        try:
            text = raw.decode(enc)
            break
        except UnicodeDecodeError:
            continue
    else:
        raise ValueError('CSV 编码无法识别(支持 UTF-8 / GBK)')
    rows = [r for r in csv.reader(text.splitlines()) if any(c.strip() for c in r)]
    if not rows:
        return [], []
    return [c.strip() for c in rows[0]], [[c.strip() for c in r] for r in rows[1:]]


def _read_excel(path) -> tuple[list[str], list[list[str]]]:
    try:
        from openpyxl import load_workbook
    except ImportError as e:  # 延迟导入:纯 CSV 用户无需 openpyxl
        raise ValueError('读取 Excel 需要安装 openpyxl') from e
    ws = load_workbook(path, read_only=True, data_only=True).active
    rows = [['' if v is None else str(v).strip() for v in row]
            for row in ws.iter_rows(values_only=True)]
    rows = [r for r in rows if any(c for c in r)]
    if not rows:
        return [], []
    return rows[0], rows[1:]
