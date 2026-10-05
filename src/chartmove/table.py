"""表格读取与取列:CSV(UTF-8/GBK)与 Excel → 图表数据,CLI --file 与 MCP file 共用。"""
from __future__ import annotations

import csv
from pathlib import Path

__all__ = [
    'read_table', 'to_float', 'parse_float', 'pick_column', 'file_columns',
    'file_box_groups', 'file_xy', 'file_samples', 'file_matrix',
]


def read_table(path: str | Path, sheet=None) -> tuple[list[str], list[list[str]]]:
    """读取表格,返回 (表头, 数据行);空行剔除;sheet 仅对 Excel 有效(名称或 1 起序号)。"""
    p = Path(path)
    suffix = p.suffix.lower()
    if suffix not in ('.csv', '.xlsx', '.xlsm'):
        raise ValueError(f'不支持的表格格式 {suffix!r}(支持 .csv / .xlsx)')
    if not p.exists():
        raise ValueError(f'文件不存在: {path}')
    if sheet is not None and suffix == '.csv':
        raise ValueError('CSV 没有工作表概念,sheet 仅对 Excel(.xlsx)有效')
    header, rows = _read_csv(p) if suffix == '.csv' else _read_excel(p, sheet)
    if not header:
        raise ValueError('表格为空')
    return header, rows


def parse_float(text, where: str = '') -> float:
    """内联标量转数值(容忍千分位逗号);nan/inf 在此合法,由校验层报中文错。"""
    s = str(text).replace(',', '').strip()
    try:
        return float(s)
    except ValueError:
        raise ValueError(f'需要数值,收到 {text!r}{where}') from None


def to_float(cell, where: str = '') -> float:
    """单元格转数值:容忍千分位逗号与空白;失败时报出位置便于定位。"""
    text = str(cell).replace(',', '').strip()
    try:
        return float(text)
    except ValueError as e:
        raise ValueError(f'数值列出现非数值单元格: {cell!r}{where}') from e


def pick_column(spec: str | None, header: list[str], default: int | None = None) -> int | None:
    """按表头名或从 1 数的序号定位列;spec 为空时返回 default。"""
    if spec is None:
        return default
    if spec in header:
        return header.index(spec)
    try:
        idx = int(spec) - 1
    except ValueError:
        raise ValueError(f'找不到列 {spec!r};可用列:{", ".join(header)}') from None
    if not 0 <= idx < len(header):
        raise ValueError(f'列序号 {spec} 超出范围,共 {len(header)} 列')
    return idx


def file_columns(path, cat_col=None, col=None,
                 sheet=None) -> tuple[list[str], list[tuple[str, list[float]]]]:
    """类目型取数:类目列(默认第 1 列)+ 数值列(默认其余全部,col 挑 1 列),
    返回 (类目, [(列名, 数值列表)]),数值列空单元格剔除。"""
    header, rows = read_table(path, sheet)
    cat_i = pick_column(cat_col, header, 0)
    cats = [r[cat_i] if len(r) > cat_i else '' for r in rows]
    val_is = [pick_column(col, header)] if col else \
        [j for j in range(len(header)) if j != cat_i]
    cols = []
    for j in val_is:
        nm = header[j] or f'列{j + 1}'
        cols.append((nm, [to_float(r[j], f'(列「{nm}」)')
                          for r in rows if len(r) > j and r[j] != '']))
    return cats, cols


def file_box_groups(path, sheet=None) -> list[tuple[str, list[float]]]:
    """箱线图取数:每列一组,首行表头 = 组名,单元格 = 原始样本(各组无需对齐)。"""
    header, rows = read_table(path, sheet)
    return [(name or f'组{j + 1}',
             [to_float(r[j], f'(列「{name}」)')
              for r in rows if len(r) > j and r[j] != ''])
            for j, name in enumerate(header)]


def file_xy(path, *, sizes=False, sheet=None) -> tuple[list[float], list[float],
                                                       list[float] | None, list[str] | None]:
    """散点/气泡取数:列顺序 x, y(, sizes)(, labels),各列独立剔除空单元格。"""
    header, rows = read_table(path, sheet)
    xs = [to_float(r[0], '(列 1)') for r in rows if r and r[0] != '']
    ys = [to_float(r[1], '(列 2)') for r in rows if len(r) > 1 and r[1] != '']
    sz = [to_float(r[2], '(列 3)') for r in rows if len(r) > 2 and r[2] != ''] \
        if sizes else None
    labels = [r[3] for r in rows if len(r) > 3 and r[3]] or None
    return xs, ys, sz, labels


def file_samples(path, sheet=None) -> list[float]:
    """直方图取数:第 1 列为原始样本。"""
    header, rows = read_table(path, sheet)
    vals = [to_float(r[0], f'(列「{header[0]}」)') for r in rows if r and r[0] != '']
    if not vals:
        raise ValueError('第 1 列需为数值样本')
    return vals


def file_matrix(path, cat_col=None, sheet=None) -> tuple[list[str], list[str], list[list[float]]]:
    """热力图取数:类目列(默认第 1 列)= 行名,其余列为矩阵(表头 = 列名)。"""
    header, rows = read_table(path, sheet)
    cat_i = pick_column(cat_col, header, 0)
    rlabels = [r[cat_i] if len(r) > cat_i else '' for r in rows]
    val_is = [j for j in range(len(header)) if j != cat_i]
    cols = [header[j] or f'列{j + 1}' for j in val_is]
    matrix = [[to_float(r[j], f'(列「{header[j]}」)') for j in val_is] for r in rows]
    return rlabels, cols, matrix


def _read_csv(path) -> tuple[list[str], list[list[str]]]:
    raw = Path(path).read_bytes()
    for enc in ('utf-8-sig', 'gbk'):
        try:
            text = raw.decode(enc)
            break
        except UnicodeDecodeError:
            continue
    else:
        raise ValueError('CSV 编码无法识别(支持 UTF-8 / GBK)')
    rows = [r for r in csv.reader(text.splitlines()) if any(c.strip() for c in r)]
    if not rows:
        return [], []
    return [c.strip() for c in rows[0]], [[c.strip() for c in r] for r in rows[1:]]


def _read_excel(path, sheet=None) -> tuple[list[str], list[list[str]]]:
    try:
        from openpyxl import load_workbook
    except ImportError as e:  # 延迟导入:纯 CSV 路径不付 openpyxl 的导入成本(包声明为硬依赖)
        raise ValueError('读取 Excel 需要安装 openpyxl') from e
    wb = load_workbook(path, read_only=True, data_only=True)
    ws = wb.active if sheet is None else wb[_resolve_sheet(wb, sheet)]
    rows = [['' if v is None else str(v).strip() for v in row]
            for row in ws.iter_rows(values_only=True)]
    rows = [r for r in rows if any(c for c in r)]
    if not rows:
        return [], []
    return rows[0], rows[1:]


def _resolve_sheet(wb, sheet) -> str:
    """按名称或从 1 数的序号定位工作表;名称优先('2024' 是常见年份表名)。"""
    names = wb.sheetnames
    if isinstance(sheet, str) and sheet in names:
        return sheet
    if isinstance(sheet, int) or (isinstance(sheet, str) and sheet.isdigit()):
        idx = int(sheet)
        if not 1 <= idx <= len(names):
            raise ValueError(f'工作表序号 {idx} 超出范围,共 {len(names)} 个;'
                             f'可用:{", ".join(names)}')
        return names[idx - 1]
    raise ValueError(f'找不到工作表 {sheet!r};可用:{", ".join(names)}')
